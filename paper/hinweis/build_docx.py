import re, sys, copy
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import content as C

TEMPLATE, OUT = sys.argv[1], sys.argv[2]
BIB = sys.argv[3]  # paper_8p_template.tex (for reference strings)
d = docx.Document(TEMPLATE)
paras = d.paragraphs

def set_text(p, text, run_idx=0):
    runs = p.runs
    runs[run_idx].text = text
    for r in runs[run_idx + 1:]:
        r._r.getparent().remove(r._r)

# ---- title block (reuse template paragraphs so their formatting is kept)
set_text(paras[4], C.TITLE)
for r in paras[4].runs: r.font.size = Pt(20)
set_text(paras[6], C.AUTHORS)
set_text(paras[7], C.AFFIL[0]); set_text(paras[8], C.AFFIL[1])
for _p in (paras[7], paras[8]):
    for r in _p.runs:
        r.font.superscript = False; r.font.subscript = False; r.font.size = Pt(9); r.font.name = "Times New Roman"
for p in (paras[9], paras[10]): p._p.getparent().remove(p._p)
abs_p = paras[13]
abs_p.runs[0].text = "Abstract"
abs_p.runs[1].text = "\u2014" + C.ABSTRACT
for r in abs_p.runs[2:]: r._r.getparent().remove(r._r)
kw = paras[15]
kw.runs[1].text = "\u2014" + C.KEYWORDS
for r in kw.runs[2:]: r._r.getparent().remove(r._r)

# ---- drop the template's empty leading table and blank paragraphs (keeps the top margin exact)
for _p in paras[0:4] + [paras[5]]:
    if not _p.text.strip() and not _p._p.xpath('.//w:sectPr'): _p._p.getparent().remove(_p._p)
for _t in d.tables[:1]:
    if not any(c.text.strip() for r in _t.rows for c in r.cells): _t._tbl.getparent().remove(_t._tbl)

# ---- remove template body after index terms (paragraph 16 onwards, all tables), keep final sectPr
body = d.element.body
keep_until = paras[15]._p
seen = False
for child in list(body):
    if child is keep_until:
        seen = True; continue
    if not seen: continue
    if child.tag == qn("w:sectPr"): continue
    body.remove(child)
final = body.find(qn("w:sectPr"))
for fr in final.findall(qn("w:footerReference")): final.remove(fr)   # no page numbers
for fr in final.findall(qn("w:headerReference")): final.remove(fr)

# ---- references: clean LaTeX bib entries
tex = open(BIB).read()
bib = {}
for m in re.finditer(r"\\bibitem\{([^}]*)\}\n(.*?)(?=\n\n\\bibitem|\n\n\\end\{thebibliography\})", tex[tex.index(r"\begin{thebibliography}"):], flags=re.S):
    s = m.group(2).replace("\n", " ")
    s = re.sub(r"\\emph\{([^}]*)\}", r"\1", s)
    s = s.replace("``", "\u201c").replace("''", "\u201d").replace("~", " ").replace("--", "\u2013").replace("\\&", "&").replace("\\_", "_")
    for a, b in (('\\"{o}', "\u00f6"), ("\\'{c}", "\u0107")): s = s.replace(a, b)
    s = re.sub(r"\s+", " ", s).strip()
    bib[m.group(1)] = s
for _k, _first in (("argyle2023", "L. P. Argyle et al."), ("shankar2026", "H. Shankar et al."), ("cao2025", "Y. Cao et al."), ("wvs7", "C. Haerpfer et al. (eds.),")):
    if _k in bib:
        _rest = bib[_k]
        _i = _rest.index("\u201c") if "\u201c" in _rest else _rest.index(",")+1
        bib[_k] = (_first + (", " if "\u201c" in _rest else " ") + _rest[_i:]).replace(" ,", ",") if "\u201c" in _rest else _first + " " + _rest[_rest.index("World Values"):]

# ---- citation numbering by first appearance
order = []
def scan(t):
    for g in re.findall(r"\[@([^\]]*)\]", t):
        for k in g.split(","):
            k = k.strip().lstrip("@")
            if k not in order: order.append(k)
for b in C.B:
    if b[0] in ("p",): scan(b[1])
    if b[0] == "table": scan(b[1])
num = {k: i + 1 for i, k in enumerate(order)}
def cite_sub(t):
    return re.sub(r"\[@([^\]]*)\]", lambda m: ", ".join("[%d]" % num[k.strip().lstrip("@")] for k in m.group(1).split(",")), t)

FONT = "Times New Roman"
def add_runs(p, text, size=9, bold=False, italic=False, smallcaps=False):
    text = cite_sub(text)
    for tok in re.split(r"(\*\*.+?\*\*|\*.+?\*)", text):
        if not tok: continue
        b, i, t = bold, italic, tok
        if tok.startswith("**"): b, t = True, tok[2:-2]
        elif tok.startswith("*") and len(tok) > 2: i, t = True, tok[1:-1]
        r = p.add_run(t); r.font.name = FONT; r.font.size = Pt(size); r.bold = b; r.italic = i; r.font.small_caps = smallcaps
        r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)

def body_par(text):
    p = d.add_paragraph(style="Body Text")
    p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.space_after = Pt(0)
    add_runs(p, text); return p

def sec_head(text):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(6); p.paragraph_format.keep_with_next = True
    add_runs(p, text, size=10, smallcaps=True); return p

def sub_head(text):
    p = d.add_paragraph(style="Heading 2"); p.paragraph_format.keep_with_next = True
    r = p.add_run(text); r.font.name = FONT; r.font.size = Pt(10); r.italic = True
    return p

def borders(cell, top=None, bottom=None):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
    for side, val in (("top", top), ("bottom", bottom)):
        if val:
            e = OxmlElement("w:" + side); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(val)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000"); b.append(e)
    tcPr.append(b)

def table_block(cap, rows, widths):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4); p.paragraph_format.keep_with_next = True
    m = re.match(r"(Table [IVX]+\.)\s*(.*)", cap)
    add_runs(p, m.group(1) + " ", size=8, smallcaps=True); add_runs(p, m.group(2), size=8, smallcaps=True)
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci); c.width = Cm(widths[ci])
            cp = c.paragraphs[0]; cp.style = d.styles["table copy"]
            cp.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = cp.add_run(val); r.font.name = FONT; r.font.size = Pt(8); r.bold = (ri == 0)
            borders(c, top=12 if ri == 0 else None, bottom=(12 if ri in (0, len(rows) - 1) else None))
    sp = d.add_paragraph(); sp.paragraph_format.space_after = Pt(2)
    sp.paragraph_format.line_spacing = Pt(4)

def fig_block(path, cap, width):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(8); p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(width))
    c = d.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after = Pt(8)
    add_runs(c, cap, size=8)

for blk in C.B:
    k = blk[0]
    if k == "h1": sec_head(blk[1])
    elif k == "h2": sub_head(blk[1])
    elif k == "p": body_par(blk[1])
    elif k == "table": table_block(blk[1], blk[2], blk[3])
    elif k == "fig": fig_block(blk[1], blk[2], blk[3])

# acknowledgment + references
sec_head("Acknowledgment")
body_par(C.ACK)
sec_head("References")
for k in order:
    p = d.add_paragraph(style="references")
    r = p.add_run(bib[k]); r.font.name = FONT; r.font.size = Pt(9)

d.core_properties.title = C.TITLE
d.core_properties.author = "K. Ayyella, B. Bhalla, Y. Bhende, I. Bansal, K. Chauhan"
d.save(OUT)
words = sum(len(b[1].split()) for b in C.B if b[0] == "p") + len(C.ABSTRACT.split())
print("saved", OUT, "| body words ~", words, "| refs", len(order), "| tables", sum(1 for b in C.B if b[0]=="table"), "figs", sum(1 for b in C.B if b[0]=="fig"))
missing=[k for k in order if k not in bib]; print("missing bib:", missing)
