import sys, re, collections, docx
from docx.shared import Emu
d = docx.Document(sys.argv[1])
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
sec = d.element.body.find(W+"sectPr")
pg = sec.find(W+"pgSz"); mar = sec.find(W+"pgMar")
tw = lambda v: round(int(v)/567, 2)
print("page cm:", tw(pg.get(W+"w")), "x", tw(pg.get(W+"h")), "| margins cm top/bottom/left/right:", tw(mar.get(W+"top")), tw(mar.get(W+"bottom")), tw(mar.get(W+"left")), tw(mar.get(W+"right")), "| cols:", sec.find(W+"cols").get(W+"num"), "| footerRef:", sec.find(W+"footerReference") is not None)
print("text width cm:", round(tw(pg.get(W+"w")) - tw(mar.get(W+"left")) - tw(mar.get(W+"right")), 2))
sizes = collections.Counter(); fonts = collections.Counter(); styles = collections.Counter()
for p in d.paragraphs:
    styles[p.style.name] += 1
    for r in p.runs:
        if r.text.strip():
            sz = r.font.size.pt if r.font.size else "style"
            sizes[(p.style.name, sz)] += len(r.text.split()); fonts[r.font.name or "style"] += 1
print("font names:", dict(fonts))
for k, v in sorted(sizes.items(), key=lambda x: -x[1])[:12]: print("  ", k, v, "words")
for t in d.tables:
    ws = [c.width.cm for c in t.rows[0].cells if c.width]
    print("table cols", len(t.columns), "width cm", round(sum(ws), 2), "(limit", round(tw(pg.get(W+"w")) - tw(mar.get(W+"left")) - tw(mar.get(W+"right")), 2), ")")
for s in d.inline_shapes: print("image width cm", round(s.width.cm, 2), "height", round(s.height.cm, 2))
body = [p for p in d.paragraphs if p.style.name == "Body Text"]
print("body paras:", len(body), "| first-line indents:", collections.Counter((p.paragraph_format.first_line_indent or 0) for p in body))
print("headings (centered small caps):", [p.text for p in d.paragraphs if p.alignment == 1 and p.text[:2] in ("I.", "II", "IV", "V.", "VI", "VII", "Ac", "Re") ][:12])
print("empty paragraphs:", sum(1 for p in d.paragraphs if not p.text.strip() and not p._p.xpath('.//w:drawing')))
txt = "\n".join(p.text for p in d.paragraphs)
print("em dashes in body:", txt.count("—") - 2, "| double spaces:", len(re.findall(r"[a-z]  [a-z]", txt)), "| 'TODO/??':", txt.count("??"))
print("cited refs:", sorted({int(x) for g in re.findall(r"\[(\d+(?:, \d+)*)\]", txt) for x in g.split(", ")}))
print("n references:", sum(1 for p in d.paragraphs if p.style.name == "references"))
