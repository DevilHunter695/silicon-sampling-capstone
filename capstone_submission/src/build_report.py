import re, sys, json
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import content_report as C

OUT = sys.argv[1]; PAGES = json.load(open(sys.argv[2])) if len(sys.argv) > 2 else {}
FONT = "Times New Roman"
d = docx.Document()
st = d.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(12); st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
def setup(sec):
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7); sec.left_margin = Cm(3.0); sec.right_margin = Cm(2.5); sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
setup(d.sections[0])

def runs(p, text, size=12, bold=False, italic=False, color=None):
    for tok in re.split(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)", text):
        if not tok: continue
        b, i, t, code = bold, italic, tok, False
        if tok.startswith("**"): b, t = True, tok[2:-2]
        elif tok.startswith("`"): code, t = True, tok[1:-1]
        elif tok.startswith("*") and len(tok) > 2: i, t = True, tok[1:-1]
        r = p.add_run(t); r.bold = b; r.italic = i; r.font.size = Pt(10 if code else size); r.font.name = "Courier New" if code else FONT
        r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
        if color: r.font.color.rgb = color
def para(text="", size=12, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, line=1.5, before=0, keep=False, indent=None):
    p = d.add_paragraph(); p.alignment = align; pf = p.paragraph_format
    pf.space_after = Pt(space_after); pf.space_before = Pt(before); pf.line_spacing = line; pf.keep_with_next = keep
    if indent is not None: pf.left_indent = Cm(indent)
    if text: runs(p, text, size, bold, italic)
    return p
def page_break(): d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
def fld(p, instr):
    r = p.add_run(); r.font.size = Pt(11); r.font.name = FONT
    for typ, txt in (("begin", None), (None, instr), ("end", None)):
        if typ: e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), typ); r._r.append(e)
        else: e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._r.append(e)
def footer_pagenum(sec, fmt, start):
    sec.footer.is_linked_to_previous = False
    for p in sec.footer.paragraphs: p.text = ""
    p = sec.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; fld(p, "PAGE")
    sp = sec._sectPr; pg = sp.find(qn("w:pgNumType"))
    if pg is None: pg = OxmlElement("w:pgNumType"); sp.append(pg)
    pg.set(qn("w:fmt"), fmt); pg.set(qn("w:start"), str(start))
def no_footer(sec):
    sec.footer.is_linked_to_previous = False
    for p in sec.footer.paragraphs: p.text = ""
def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color); tcPr.append(sh)
def cell_borders(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
    for side, val in kw.items():
        e = OxmlElement("w:" + side); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(val)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000"); b.append(e)
    tcPr.append(b)

# ---------------- cover ----------------
no_footer(d.sections[0])
t = d.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
c0, c1 = t.cell(0, 0), t.cell(0, 1); c0.width = Cm(4.0); c1.width = Cm(11.5); shade(c0, "001D3E"); shade(c1, "001D3E")
c0.paragraphs[0].add_run().add_picture("logo_nmims.png", width=Cm(3.2)); c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
p = c1.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT; runs(p, "MUKESH PATEL SCHOOL OF TECHNOLOGY MANAGEMENT & ENGINEERING", 14, True, color=RGBColor(255, 255, 255))
p = c1.add_paragraph(); runs(p, "SVKM's NMIMS (Deemed-to-be University)", 11, color=RGBColor(0xdb, 0xe4, 0xff))
p = c1.add_paragraph(); runs(p, "Department of Computer Engineering", 11, color=RGBColor(0xdb, 0xe4, 0xff))
para("", space_after=18)
para("CAPSTONE PROJECT REPORT", 14, True, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0, space_after=14)
para(C.TITLE.upper(), 24, True, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.1, space_after=8)
para(C.SUBTITLE, 13, False, True, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.2, space_after=26)
para("Submitted by", 12, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0, space_after=4)
para("GROUP " + C.GROUP, 14, True, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0, space_after=8)
tb = d.add_table(rows=len(C.STUDENTS) + 1, cols=2); tb.alignment = WD_TABLE_ALIGNMENT.CENTER; tb.autofit = False
for ri, (nm, sap) in enumerate([("Student name", "SAP ID")] + C.STUDENTS):
    for ci, v in enumerate((nm, sap)):
        c = tb.cell(ri, ci); c.width = Cm(6.5 if ci == 0 else 4.5); cp = c.paragraphs[0]; cp.alignment = WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_after = Pt(2)
        runs(cp, v, 12, ri == 0); cell_borders(c, top=4, bottom=4, left=4, right=4)
        if ri == 0: shade(c, "E7ECF7")
para("", space_after=10)
para("Under the guidance of", 12, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0, space_after=4)
para(C.MENTOR, 14, True, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0, space_after=22)
para("Academic Year 2026-2027", 12, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0, space_after=2)
para("October 2026", 12, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0, space_after=2)

# ---------------- front matter (roman) ----------------
sec = d.add_section(WD_SECTION.NEW_PAGE); setup(sec); footer_pagenum(sec, "lowerRoman", 1)
def front_title(t):
    para(t, 14, True, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0, space_after=16, before=6)
front_title("CERTIFICATE")
para("This is to certify that the capstone project entitled **“" + C.TITLE + "”** is a bonafide work carried out by the following students of Group " + C.GROUP + ", Department of Computer Engineering, Mukesh Patel School of Technology Management & Engineering, NMIMS, during the academic year 2026-2027, under my guidance, in partial fulfilment of the requirements of the Capstone Project.")
for nm, sap in C.STUDENTS: para(f"{nm}  (SAP ID: {sap})", 12, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.3, space_after=2, indent=1.0)
para("", space_after=36)
tb = d.add_table(rows=2, cols=2); tb.autofit = False
for ci, (nm, role) in enumerate(((C.MENTOR, "Project Mentor"), ("____________________", "Head of Department / Capstone Committee"))):
    c = tb.cell(0, ci); c.width = Cm(7.7); runs(c.paragraphs[0], "\n\n_______________________", 12)
    c2 = tb.cell(1, ci); c2.width = Cm(7.7); runs(c2.paragraphs[0], nm if ci == 0 else "", 12, True); p2 = c2.add_paragraph(); runs(p2, role, 11)
para("", space_after=8)
para("Date: ____________________          Place: ____________________", 12, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.3)
page_break()
front_title("DECLARATION")
para("We declare that this report, **“" + C.TITLE + "”**, describes the work carried out by us as part of the capstone project at the Mukesh Patel School of Technology Management & Engineering, NMIMS, under the guidance of " + C.MENTOR + ". All sources of data, ideas and text taken from the work of others have been acknowledged and cited in the references. The surveys used are public datasets and are used in line with the terms of their publishers; respondent-level microdata are not redistributed.")
para("", space_after=26)
tb = d.add_table(rows=len(C.STUDENTS), cols=2); tb.autofit = False
for ri, (nm, sap) in enumerate(C.STUDENTS):
    c = tb.cell(ri, 0); c.width = Cm(8.0); runs(c.paragraphs[0], f"{nm} ({sap})", 12); c.paragraphs[0].paragraph_format.space_after = Pt(18)
    c = tb.cell(ri, 1); c.width = Cm(7.5); runs(c.paragraphs[0], "Signature: ____________________", 12)
page_break()
front_title("ACKNOWLEDGEMENT"); para(C.ACK)
page_break()
front_title("ABSTRACT"); para(C.ABSTRACT, line=1.4); para("**Keywords:** " + C.KEYWORDS, line=1.4, before=6)
page_break()

# contents
front_title("CONTENTS")
toc = []
for b in C.B:
    if b[0] == "h1": toc.append((1, b[1]))
    if b[0] == "h2": toc.append((2, b[1]))
toc.append((1, "REFERENCES"))
for tt, _ in C.APPX: toc.append((1, "APPENDIX " + tt))
for lvl, name in toc:
    p = d.add_paragraph(); pf = p.paragraph_format; pf.line_spacing = 1.15; pf.space_after = Pt(2 if lvl == 2 else 4); pf.left_indent = Cm(0 if lvl == 1 else 0.8)
    pf.tab_stops.add_tab_stop(Cm(15.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    runs(p, name + "\t" + str(PAGES.get(name, "0")), 11 if lvl == 2 else 12, lvl == 1)

# ---------------- main body (arabic) ----------------
sec = d.add_section(WD_SECTION.NEW_PAGE); setup(sec); footer_pagenum(sec, "decimal", 1)
fig_n = [0]
def h1(t, first=False):
    if not first: page_break()
    para(t, 14, True, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.0, space_after=12, keep=True)
def h2(t): para(t, 12, True, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.2, space_after=4, before=6, keep=True)
def bullets(items):
    for it in items:
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; pf = p.paragraph_format; pf.left_indent = Cm(0.9); pf.first_line_indent = Cm(-0.5); pf.line_spacing = 1.4; pf.space_after = Pt(3)
        runs(p, "•  " + it, 12)
def table_block(cap, rows, widths):
    p = para(cap, 11, True, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.1, space_after=4, before=6, keep=True)
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for ri, row in enumerate(rows):
        for ci, v in enumerate(row):
            c = t.cell(ri, ci); c.width = Cm(widths[ci]); cp = c.paragraphs[0]; cp.paragraph_format.space_after = Pt(1); cp.paragraph_format.line_spacing = 1.05
            cp.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 or len(v) > 22 else WD_ALIGN_PARAGRAPH.CENTER
            runs(cp, v, 10, ri == 0)
            cell_borders(c, top=4, bottom=4, left=4, right=4)
            if ri == 0: shade(c, "E7ECF7")
            if ri < len(rows) - 1: cp.paragraph_format.keep_with_next = True
    para("", space_after=4, line=1.0)
def fig_block(path, cap, w):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(path, width=Cm(w))
    para(cap, 11, False, True, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.1, space_after=10)
def box_block(cap, text):
    para(cap, 11, True, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.1, space_after=4, before=6, keep=True)
    t = d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    c = t.cell(0, 0); c.width = Cm(14.5); cell_borders(c, top=6, bottom=6, left=6, right=6); shade(c, "F4F6FB")
    first = True
    for line in text.split("\n"):
        cp = c.paragraphs[0] if first else c.add_paragraph(); first = False
        cp.paragraph_format.space_after = Pt(0); cp.paragraph_format.line_spacing = 1.1; runs(cp, line, 10)
    para("", space_after=6, line=1.0)
def render(blocks):
    for b in blocks:
        k = b[0]
        if k == "p": para(b[1])
        elif k == "bul": bullets(b[1])
        elif k == "table": table_block(b[1], b[2], b[3])
        elif k == "fig": fig_block(b[1], b[2], b[3])
        elif k == "box": box_block(b[1], b[2])
first = True
for b in C.B:
    if b[0] == "h1": h1(b[1], first); first = False
    elif b[0] == "h2": h2(b[1])
    else: render([b])
h1("REFERENCES")
for i, r in enumerate(C.REFS, 1):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; pf = p.paragraph_format; pf.left_indent = Cm(1.0); pf.first_line_indent = Cm(-1.0); pf.line_spacing = 1.2; pf.space_after = Pt(5)
    runs(p, f"[{i}]\t" + r, 11); pf.tab_stops.add_tab_stop(Cm(1.0))
for tt, blocks in C.APPX:
    h1("APPENDIX " + tt); render(blocks)
d.core_properties.title = C.TITLE; d.core_properties.author = "Group 55 (B. Bhalla, I. Bansal, Y. Bhende, K. Ayyella)"
d.save(OUT); print("saved", OUT)
json.dump([n for _, n in toc], open("toc_names.json", "w"))
